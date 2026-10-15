"""Build the CPLH-DRIVEN warehouse payroll forecast model (synthetic data) as an Excel workbook.

Built from the simple model in simple/, with the chain turned around: CPLH and cases are the
inputs, and payroll (and payroll % of sales) is the result.

    Cases / CPLH = hours;  hours x cost per hour = staff payroll;
    staff payroll / (1 - leadership share) = total payroll;  total payroll / sales = payroll %.

CPLH, cases per $ and cost per hour are the same month last year (CPLH and cases can be adjusted
or overridden on their own tabs). The leadership share comes from the last actual month.

Run (from the repo root):  .venv/bin/python cplh-driven/build_model.py
Output: cplh-driven/warehouse_payroll_forecast_cplh.xlsx
"""
import calendar
import os
import random
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "warehouse_payroll_forecast_cplh.xlsx")
# bounded ranges (whole-column refs are slow); room for pasted data up to these rows
DATA_MAX, HOL_MAX = 5000, 1000
DR = lambda col: f"Data_Actuals!${col}$2:${col}${DATA_MAX}"
PR = lambda col: f"Plan_Inputs!${col}$2:${col}${DATA_MAX}"
HR = lambda col: f"Holidays!${col}$2:${col}${HOL_MAX}"
START = date(2025, 1, 1)
END = date(2028, 12, 1)
LAST_ACTUAL = date(2026, 9, 1)
MONTHS = []
d = START
while d <= END:
    MONTHS.append(d)
    d = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
N = len(MONTHS)  # 48
FIRST_ROW = 15
LAST_ROW = FIRST_ROW + N - 1  # 62
YEARS = [2025, 2026, 2027, 2028]

# code, type, state, host region (IDCs borrow host sales)
SITES = [
    ("NSW-1", "Region", "NSW", None),
    ("VIC-1", "Region", "VIC", None),
    ("QLD-1", "Region", "QLD", None),
    ("NSW-2", "Region", "NSW", None),
    ("VIC-2", "Region", "VIC", None),
    ("QLD-2", "Region", "QLD", None),
    ("SA-1", "Region", "SA", None),
    ("WA-1", "Region", "WA", None),
    ("IDC-NSW", "IDC", "NSW", "NSW-1"),
    ("IDC-VIC", "IDC", "VIC", "VIC-1"),
]
REGIONS = [s[0] for s in SITES if s[1] == "Region"]
# regions each IDC ships to: their sales drive IDC volume (payroll % stays on host sales)
SERVED = {"IDC-NSW": ["NSW-1", "NSW-2"], "IDC-VIC": ["VIC-1", "VIC-2"]}
STATES = ["NSW", "VIC", "QLD", "SA", "WA"]
# every site code in the workbook reads from the Code column of the Sites table on Inputs,
# so typing new codes there renames the sites everywhere (tab names are cosmetic)
SITE_ROW = {s[0]: 36 + i for i, s in enumerate(SITES)}
SITE_IDX = {s[0]: i for i, s in enumerate(SITES)}
CODE = lambda c: f"Inputs!$A${SITE_ROW[c]}"

WAGE_TABLE = [(date(2026, 1, 1), 0.040), (date(2027, 1, 1), 0.045), (date(2028, 1, 1), 0.045)]
PH_ROSTER = 0.5   # share of a normal day's hours worked on a public holiday
PH_LOADING = 1.5  # extra pay on a public holiday (1.5 = paid 250%)


# ---------------------------------------------------------------- holidays
def easter(y):
    a, b, c = y % 19, y // 100, y % 100
    dd, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - dd - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    month = (h + l - 7 * m + 114) // 31
    day = (h + l - 7 * m + 114) % 31 + 1
    return date(y, month, day)


def nth_weekday(y, m, wd, n):
    """n-th weekday (0=Mon) of a month; n=-1 for the last."""
    if n > 0:
        d0 = date(y, m, 1)
        return d0 + timedelta(days=(wd - d0.weekday()) % 7 + 7 * (n - 1))
    d0 = date(y, m, calendar.monthrange(y, m)[1])
    return d0 - timedelta(days=(d0.weekday() - wd) % 7)


def next_monday(d0):
    return d0 + timedelta(days=(7 - d0.weekday()) % 7)


def holidays():
    rows = []
    for y in YEARS:
        ny = date(y, 1, 1)
        aus = date(y, 1, 26)
        gf = easter(y) - timedelta(days=2)
        xmas = date(y, 12, 25)
        anzac = date(y, 4, 25)
        for st in STATES:
            h = [(ny, "New Year's Day")]
            if ny.weekday() >= 5:
                h.append((next_monday(ny), "New Year's Day (additional)"))
            h.append((next_monday(aus) if aus.weekday() >= 5 else aus, "Australia Day"))
            h.append((gf, "Good Friday"))
            if st != "WA":
                h.append((gf + timedelta(days=1), "Easter Saturday"))
            h.append((gf + timedelta(days=2), "Easter Sunday"))
            h.append((gf + timedelta(days=3), "Easter Monday"))
            h.append((anzac, "Anzac Day"))
            if (st == "WA" and anzac.weekday() >= 5) or (st == "SA" and anzac.weekday() == 6):
                h.append((next_monday(anzac), "Anzac Day (additional)"))
            if st in ("NSW", "VIC", "SA"):
                h.append((nth_weekday(y, 6, 0, 2), "King's Birthday"))
            elif st == "QLD":
                h.append((nth_weekday(y, 10, 0, 1), "King's Birthday"))
            else:
                h.append((nth_weekday(y, 9, 0, -1), "King's Birthday"))
            labour = {"NSW": (10, 1), "SA": (10, 1), "VIC": (3, 2), "QLD": (5, 1), "WA": (3, 1)}[st]
            h.append((nth_weekday(y, labour[0], 0, labour[1]), "Labour Day"))
            if st == "VIC":
                h.append((nth_weekday(y, 11, 1, 1), "Melbourne Cup"))
                h.append((nth_weekday(y, 9, 5, -1) - timedelta(days=1), "AFL Grand Final Friday (check date)"))
            if st == "SA":
                h.append((nth_weekday(y, 3, 0, 2), "Adelaide Cup"))
            if st == "WA":
                h.append((nth_weekday(y, 6, 0, 1), "WA Day"))
            h.append((xmas, "Christmas Day"))
            h.append((date(y, 12, 26), "Proclamation Day" if st == "SA" else "Boxing Day"))
            if xmas.weekday() == 5:
                h += [(date(y, 12, 27), "Christmas (additional)"), (date(y, 12, 28), "Boxing Day (additional)")]
            elif xmas.weekday() == 6:
                h.append((date(y, 12, 27), "Christmas (additional)"))
            elif xmas.weekday() == 4:
                h.append((date(y, 12, 28), "Boxing Day (additional)"))
            rows += [(dt, name, st) for dt, name in h]
    rows.sort(key=lambda r: (r[2], r[0]))
    return rows


HOLIDAYS = holidays()


def ph_factor(month, state):
    days = calendar.monthrange(month.year, month.month)[1]
    h = sum(1 for dt, _, st in HOLIDAYS if st == state and dt.year == month.year and dt.month == month.month)
    return (days - h + h * PH_ROSTER * (1 + PH_LOADING)) / (days - h + h * PH_ROSTER)


def wage_index(month):
    idx = 1.0
    for eff, pct in WAGE_TABLE:
        if eff <= month:
            idx *= 1 + pct
    return idx


# ---------------------------------------------------------- synthetic data
SEASON_SALES = {1: .97, 2: .90, 3: 1.0, 4: .99, 5: .98, 6: .96, 7: .98, 8: .97, 9: .96, 10: 1.0, 11: 1.03, 12: 1.25}
SEASON_CPLH = {1: .97, 2: .99, 3: 1.0, 4: .99, 5: 1.0, 6: 1.0, 7: 1.0, 8: 1.0, 9: 1.0, 10: 1.01, 11: 1.03, 12: 1.06}
SEASON_LEAD = {11: 1.05, 12: 1.15}
STATE_CPH = {"NSW": 66.0, "VIC": 65.0, "QLD": 63.0, "SA": 61.0, "WA": 69.0}
# monthly sales $m (2025), cases per $ sales, CPLH, leadership $m per month
PARAMS = {
    "NSW-1": (150, .0210, 78, .32), "VIC-1": (140, .0215, 80, .30), "QLD-1": (115, .0220, 74, .26),
    "NSW-2": (125, .0212, 76, .28), "VIC-2": (135, .0208, 79, .30), "QLD-2": (120, .0218, 75, .27),
    "SA-1": (80, .0225, 72, .21), "WA-1": (105, .0223, 73, .25),
    "IDC-NSW": (None, .0030, 110, .15), "IDC-VIC": (None, .0028, 105, .14),
}


def synth():
    rng = random.Random(42)
    data = {}
    for code, typ, state, host in SITES:
        base_sales, cps0, cplh0, lead0 = PARAMS[code]
        for t, m in enumerate(MONTHS):
            if typ == "Region":
                sales = base_sales * 1e6 * SEASON_SALES[m.month] * 1.035 ** (t / 12) * (1 + rng.gauss(0, .015))
            else:
                sales = data[(host, m)]["sales"]
            cps = cps0 * 0.985 ** (t / 12) * (0.98 if m.month == 12 else 1) * (1 + rng.gauss(0, .01))
            driver = sum(data[(s, m)]["sales"] for s in SERVED[code]) if typ == "IDC" else sales
            cases = driver * cps
            cplh = cplh0 * SEASON_CPLH[m.month] * 1.01 ** (t / 12) * (1 + rng.gauss(0, .015))
            hours = cases / cplh
            cph = STATE_CPH[state] * (1.06 if typ == "IDC" else 1) * wage_index(m) * ph_factor(m, state) * (1 + rng.gauss(0, .008))
            staff = hours * cph
            lead = lead0 * 1e6 * wage_index(m) * SEASON_LEAD.get(m.month, 1) * (1 + rng.gauss(0, .02))
            data[(code, m)] = dict(sales=sales, cases=cases, hours=hours, staff=staff, lead=lead)
    return data


DATA = synth()

# ------------------------------------------------------------------ styles
NAVY = PatternFill("solid", fgColor="1F3864")
LIGHT = PatternFill("solid", fgColor="DDEBF7")
INPUT = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="F2F2F2")
WHITE_B = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
BLUE = Font(color="0000FF")
HELP = Font(color="808080", italic=True)
THIN = Side(style="thin", color="1F3864")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

F_M = '#,##0.0,,'        # $ or units shown in millions
F_CASES = '#,##0.00,,'
F_INT = '#,##0'
F_1 = '0.0'
F_CPH = '$#,##0.00'
F_CPC = '$#,##0.000'
F_PCT = '0.00%'
F_DATE = 'mmm-yy'
F_RATIO = '0.00000'

wb = Workbook()


def header(ws, row, col, text, fill=NAVY, font=WHITE_B, width=None):
    c = ws.cell(row=row, column=col, value=text)
    c.fill, c.font, c.alignment = fill, font, CENTER
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return c


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)




# ------------------------------------------------------------------ README
ws = wb.active
ws.title = "README"
readme = [
    ("Warehouse payroll forecast model (CPLH-driven version)", BOLD),
    ("Synthetic data only. Monthly, Jan 2025 to Dec 2028.", None),
    ("", None),
    ("THE MODEL (per site, per forecast month)", BOLD),
    ("CPLH and cases are the inputs. Payroll, and payroll % of sales, are the result.", None),
    ("CPLH comes from the CPLH_Forecast tab: CPLH in the same month last year, plus your adjustments, or your override.", None),
    ("Cases come from the Cases_Forecast tab: volume driver sales x cases per $ in the same month last year,", None),
    ("   plus your adjustments, or your override.", None),
    ("Hours = cases / CPLH.", None),
    ("Cost per hour = same month last year's cost per hour, with last year's public holidays and wage index taken out", None),
    ("   and this month's put back in.  Staff payroll = hours x cost per hour.", None),
    ("Total payroll = staff payroll / (1 - leadership share of total payroll in the last actual month).", None),
    ("Leadership = total payroll - staff payroll.", None),
    ("Payroll % of sales = total payroll / sales (plan sales from Plan_Inputs).  Cost per case = total payroll / cases.", None),
    ("", None),
    ("SAME MONTH LAST YEAR", BOLD),
    ("CPLH, cases per $ and cost per hour look 12 rows up. If that month is an actual, the rate comes from actuals;", None),
    ("if it is a forecast, it uses that month's run rate. So Mar-28 uses Mar-27, which uses Mar-26 actuals.", None),
    ("Adjustments and overrides are not carried into the next year. Roll forward by moving Last actual month.", None),
    ("", None),
    ("HOW TO UPDATE", BOLD),
    ("1. Paste new actuals into Data_Actuals and the plan into Plan_Inputs (same columns).", None),
    ("2. Move 'Last actual month' on Inputs forward.", None),
    ("3. Review Cases_Forecast and CPLH_Forecast; clear adjustments or overrides for months that are now actuals.", None),
    ("4. Check the Checks tab says ALL CHECKS OK.", None),
    ("", None),
    ("SITE CODES", BOLD),
    ("Type your own codes in the Code column of the Sites table on Inputs; every tab follows. Renaming tabs is optional.", None),
    ("", None),
    ("COLOURS", BOLD),
    ("Yellow cells with blue text are inputs. Blue-shaded rows are forecast months. Grey columns are helper calculations.", None),
    ("Public holidays are generated by rule; check against the official state lists (AFL Grand Final Friday in particular).", None),
]
for i, (txt, font) in enumerate(readme, 1):
    c = ws.cell(row=i, column=1, value=txt)
    if font:
        c.font = font
ws.column_dimensions["A"].width = 120


def inp(c):
    c.fill, c.font = INPUT, BLUE


# ------------------------------------------------------------------ Inputs
ws = wb.create_sheet("Inputs")
ws.column_dimensions["A"].width = 34
for col in "BCDEFG":
    ws.column_dimensions[col].width = 14
ws["A1"] = "Inputs"
ws["A1"].font = Font(bold=True, size=14)
ws["A3"] = "Last actual month"
ws["B3"] = LAST_ACTUAL
ws["B3"].number_format = F_DATE
inp(ws["B3"])
ws["A5"] = "Model status"
ws["B5"] = "=Checks!B3"
ws["B5"].font = BOLD
name("LastActual", "Inputs!$B$3")
ws["A7"] = "Warn if forecast CPLH moves more than this vs LY"
ws["B7"] = 0.10
ws["B7"].number_format = '0.0%'
inp(ws["B7"])
name("CPLHTol", "Inputs!$B$7")

ws["A19"] = "Wage increases (staff cost per hour)"
ws["A19"].font = BOLD
header(ws, 20, 1, "Effective month")
header(ws, 20, 2, "Increase")
for i in range(8):
    r = 21 + i
    a, b = ws.cell(row=r, column=1), ws.cell(row=r, column=2)
    if i < len(WAGE_TABLE):
        a.value, b.value = WAGE_TABLE[i]
    a.number_format, b.number_format = F_DATE, '0.0%'
    inp(a)
    inp(b)
ws["C21"] = "Include past increases too: last year's cost per hour is divided by its wage index before this year's is applied."
ws["C21"].font = HELP
name("WageDates", "Inputs!$A$21:$A$28")
name("WageRates", "Inputs!$B$21:$B$28")

ws["A30"] = "Public holidays"
ws["A30"].font = BOLD
ws["A31"] = "Share of normal hours worked on a PH"
ws["B31"] = PH_ROSTER
ws["A32"] = "PH pay loading (1.5 = paid 250%)"
ws["B32"] = PH_LOADING
inp(ws["B31"])
inp(ws["B32"])
name("PH_Roster", "Inputs!$B$31")
name("PH_Loading", "Inputs!$B$32")

ws["A34"] = "Sites"
ws["A34"].font = BOLD
ws["B34"] = "Type your own codes in the Code column; every tab follows. Renaming the tabs to match is optional."
ws["B34"].font = HELP
for j, t in enumerate(["Code", "Type", "State", "Payroll % basis", "In national sales?", "Volume driver"], 1):
    header(ws, 35, j, t)
for i, (code, typ, st, host) in enumerate(SITES):
    r = 36 + i
    vals = [code, typ, st, f"=A{SITE_ROW[host or code]}", "No" if typ == "IDC" else "Yes",
            "=" + '&"+"&'.join(f"A{SITE_ROW[s]}" for s in SERVED.get(code, [code]))]
    for j, v in enumerate(vals, 1):
        ws.cell(row=r, column=j, value=v)
    inp(ws.cell(row=r, column=1))
name("SiteCodes", f"Inputs!$A$36:$A${35 + len(SITES)}")

# ------------------------------------------------------------ Data_Actuals
ws = wb.create_sheet("Data_Actuals")
cols = ["Month", "Site", "Sales $", "Cases picked", "WH hours worked", "Staff payroll $", "Leadership payroll $"]
for j, t in enumerate(cols, 1):
    header(ws, 1, j, t, width=16)
r = 2
for m in MONTHS:
    if m > LAST_ACTUAL:
        break
    for code, typ, _, _ in SITES:
        v = DATA[(code, m)]
        row = [m, "=" + CODE(code), None if typ == "IDC" else round(v["sales"]), round(v["cases"]), round(v["hours"]),
               round(v["staff"]), round(v["lead"])]
        for j, val in enumerate(row, 1):
            c = ws.cell(row=r, column=j, value=val)
            c.number_format = F_DATE if j == 1 else F_INT
        r += 1
ws.freeze_panes = "A2"

# ------------------------------------------------------------- Plan_Inputs
ws = wb.create_sheet("Plan_Inputs")
for j, t in enumerate(["Month", "Site", "Sales forecast $"], 1):
    header(ws, 1, j, t, width=22)
ws["E1"] = "From the long-range plan. IDCs have no sales; their payroll % is shown against the host region's sales."
ws["F1"].font = HELP
r = 2
for m in MONTHS:
    if m.year < 2026:
        continue
    for code, typ, _, host in SITES:
        v = DATA[(code, m)]
        row = [m, "=" + CODE(code), None if typ == "IDC" else round(v["sales"], -3)]
        for j, val in enumerate(row, 1):
            c = ws.cell(row=r, column=j, value=val)
            c.number_format = [F_DATE, None, F_INT][j - 1] or "General"
        r += 1
ws.freeze_panes = "A2"

# ------------------------------------------------- Cases_Forecast / CPLH_Forecast
# months down, sites across, in four blocks: run rate | adjustments | overrides | used
CF_FIRST = 8
CF_LAST = CF_FIRST + N - 1
cf_col = lambda block, i: get_column_letter(3 + block * (len(SITES) + 1) + i)


def forecast_tab(title, heading, unit, used_title, notes, run_rate, fmt):
    """run_rate(code, site_row) -> formula for the run-rate cell of a forecast month."""
    ws = wb.create_sheet(title)
    ws["A1"] = heading
    ws["A1"].font = Font(bold=True, size=14)
    for k, line in enumerate(notes):
        ws.cell(row=2 + k, column=1, value=line).font = HELP
    header(ws, 7, 1, "Month", width=9)
    header(ws, 7, 2, "Status", width=9)
    blocks = ["1. Run-rate forecast", f"2. Adjustments (+/- {unit})", "3. Override (replaces 1 + 2)", used_title]
    for b, t in enumerate(blocks):
        first, last = cf_col(b, 0), cf_col(b, len(SITES) - 1)
        ws.merge_cells(f"{first}6:{last}6")
        header(ws, 6, ws[f"{first}1"].column, t)
        ws.column_dimensions[get_column_letter(ws[f"{last}1"].column + 1)].width = 2
        for i, (code, *_r) in enumerate(SITES):
            header(ws, 7, ws[f"{cf_col(b, i)}1"].column, "=" + CODE(code), width=11)
    for i, m in enumerate(MONTHS):
        r, src_r = CF_FIRST + i, FIRST_ROW + i
        ws[f"A{r}"] = m
        ws[f"A{r}"].number_format = F_DATE
        ws[f"B{r}"] = f'=IF(A{r}<=LastActual,"Actual","Forecast")'
        for j, (code, *_r) in enumerate(SITES):
            rr, adj, ovr, used = (cf_col(b, j) for b in range(4))
            ws[f"{rr}{r}"] = run_rate(code, src_r, r)
            ws[f"{used}{r}"] = f'=IF($B{r}="Actual",{rr}{r},IF({ovr}{r}<>"",{ovr}{r},{rr}{r}+N({adj}{r})))'
            inp(ws[f"{adj}{r}"])
            inp(ws[f"{ovr}{r}"])
            for col in (rr, adj, ovr, used):
                ws[f"{col}{r}"].number_format = fmt
            ws[f"{used}{r}"].font = BOLD
    ws.conditional_formatting.add(f"A{CF_FIRST}:{cf_col(3, len(SITES) - 1)}{CF_LAST}",
                                  FormulaRule(formula=[f'$B{CF_FIRST}="Actual"'], fill=GREY))
    ws.freeze_panes = f"C{CF_FIRST}"


forecast_tab(
    "Cases_Forecast", "Cases forecast", "cases", "4. Cases forecast used",
    ["1. Run rate: volume driver sales x cases per $ of sales in the same month last year. Actual months show actuals.",
     "2. Adjustments: type extra cases (or negative) for a site and month, e.g. a new store or volume moving between sites.",
     "3. Override: type the cases you want for a site and month. It replaces 1 + 2. Leave blank to use the model.",
     "4. Cases used = override if there is one, otherwise run rate + adjustment. This feeds Cases picked on each site tab."],
    lambda code, sr, r: f"=IF($B{r}=\"Actual\",'{code}'!S{sr},IFERROR('{code}'!Y{sr}*'{code}'!AB{sr},0))", F_INT)
forecast_tab(
    "CPLH_Forecast", "CPLH forecast (cases per labour hour)", "cases / hr", "4. CPLH forecast used",
    ["1. Run rate: CPLH in the same month last year (actual, else that month's run rate). Actual months show actuals.",
     "2. Adjustments: type a change in cases per hour for a site and month, e.g. +2 for a new pick system.",
     "3. Override: type the CPLH you want for a site and month, e.g. a target. It replaces 1 + 2. Leave blank to use the model.",
     "4. CPLH used = override if there is one, otherwise run rate + adjustment. This drives hours, and so payroll, on each site tab."],
    # column X (actual CPLH, else same month LY), not K: K's forecast rows read this tab
    lambda code, sr, r: f"=IFERROR('{code}'!X{sr}*1,0)", F_1)
CF_RANGE = lambda tab, b: f"{tab}!${cf_col(b, 0)}${CF_FIRST}:${cf_col(b, len(SITES) - 1)}${CF_LAST}"
CF_STATUS = lambda tab: f"{tab}!$B${CF_FIRST}:$B${CF_LAST}"

# ---------------------------------------------------------------------- CLR
ws = wb.create_sheet("CLR")
ws["A1"] = "CLR: cost-saving initiatives (placeholder)"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "Placeholder only. Nothing on this tab feeds the forecast yet."
ws["A2"].font = HELP
for j, (t, wd) in enumerate([("Initiative", 30), ("Site", 10), ("Type", 18), ("Value", 11), ("Start month", 11),
                              ("Ramp-up months", 11), ("End month", 11), ("Active", 8), ("Owner / notes", 30)], 1):
    header(ws, 4, j, t, width=wd)

# ---------------------------------------------------------------- Holidays
ws = wb.create_sheet("Holidays")
for j, t in enumerate(["Date", "Holiday", "State"], 1):
    header(ws, 1, j, t, width=[12, 38, 8][j - 1])
for i, (dt, nm, st) in enumerate(HOLIDAYS, 2):
    ws.cell(row=i, column=1, value=dt).number_format = "dd-mmm-yy"
    ws.cell(row=i, column=2, value=nm)
    ws.cell(row=i, column=3, value=st)
ws["E1"] = "Generated by rule. Add, remove or move rows freely; the Calendar tab counts them by state and month."
ws["E1"].font = HELP
ws.freeze_panes = "A2"

# ---------------------------------------------------------------- Calendar
ws = wb.create_sheet("Calendar")
ws["A1"] = "Calendar: public holidays, PH cost factor, wage index"
ws["A1"].font = BOLD
ws["A2"] = "PH factor = average cost per hour in the month vs a month with no public holidays."
ws["A2"].font = HELP
heads = ["Month", "Days"] + [f"PH {s}" for s in STATES] + STATES + ["Wage index"]
for j, t in enumerate(heads, 1):
    header(ws, 4, j, t, width=11)
ws.cell(row=3, column=3, value="Public holidays (count)").font = BOLD
ws.cell(row=3, column=8, value="PH cost factor").font = BOLD
for i, m in enumerate(MONTHS):
    r = 5 + i
    ws.cell(row=r, column=1, value=m).number_format = F_DATE
    ws.cell(row=r, column=2, value=f"=DAY(EOMONTH(A{r},0))")
    for k, st in enumerate(STATES):
        cnt = 3 + k
        ws.cell(row=r, column=cnt, value=(
            f'=COUNTIFS({HR("A")},">="&$A{r},{HR("A")},"<="&EOMONTH($A{r},0),{HR("C")},"{st}")'))
        h = get_column_letter(cnt)
        c = ws.cell(row=r, column=8 + k, value=(
            f"=($B{r}-{h}{r}+{h}{r}*PH_Roster*(1+PH_Loading))/($B{r}-{h}{r}+{h}{r}*PH_Roster)"))
        c.number_format = '0.0000'
    # product of (1 + increase) for every increase effective on or before this month
    c = ws.cell(row=r, column=13, value=f"=EXP(SUMPRODUCT((WageDates<=$A{r})*(WageDates<>\"\")*LN(1+WageRates)))")
    c.number_format = '0.0000'
ws.freeze_panes = "B5"

# -------------------------------------------------------------- site tabs
SITE_COLS = {
    "A": ("Month", F_DATE, 9), "B": ("Year", "0", 6), "C": ("Mth", "0", 5), "D": ("Status", None, 9),
    "E": ("Sales", F_M, 10), "F": ("Cases picked", F_CASES, 10), "G": ("WH hours worked", F_INT, 11),
    "H": ("Staff payroll", F_M, 10), "I": ("Leadership payroll", F_M, 11), "J": ("Total payroll", F_M, 10),
    "K": ("CPLH (cases / hr)", F_1, 10), "L": ("Cost per hour (staff)", F_CPH, 10),
    "M": ("Cost per case (total payroll)", F_CPC, 11), "N": ("Total payroll % sales", F_PCT, 10),
    "O": ("Leadership % of total payroll", F_PCT, 11),
    "R": ("Act sales", F_INT, 11), "S": ("Act cases", F_INT, 11), "T": ("Act hours", F_INT, 10),
    "U": ("Act staff $", F_INT, 11), "V": ("Act leadership $", F_INT, 11),
    "W": ("Plan sales", F_INT, 12), "X": ("CPLH run rate (actual, else same month LY)", F_1, 10),
    "Y": ("Volume driver sales", F_INT, 12),
    "Z": ("PH factor", '0.0000', 8), "AA": ("Wage index", '0.0000', 8),
    "AB": ("Cases per $ (actual, else same month LY)", F_RATIO, 11),
    "AC": ("Cost per hour ex PH & wage (actual, else same month LY)", F_CPH, 12),
    "AD": ("Leadership share of total payroll (actual, else last actual month)", F_PCT, 12),
    "AE": ("Fcst cost per hour (AC x PH x wage)", F_CPH, 11),
    "AF": ("LY CPLH (same month, as used)", F_1, 9), "AG": ("Fcst CPLH vs LY", F_PCT, 9),
}
UNITS = {"E": "$m", "F": "m", "G": "#", "H": "$m", "I": "$m", "J": "$m", "K": "#", "L": "$", "M": "$",
         "N": "%", "O": "%"}
METRIC_COLS = "KLMNO"
RNG = lambda col: f"${col}${FIRST_ROW}:${col}${LAST_ROW}"


def data_sum(col, r):
    return f"SUMIFS({DR(col)},{DR('A')},$A{r},{DR('B')},$B$2)"


def metric_formulas(ws, r):
    ws[f"K{r}"] = f"=IFERROR(F{r}/G{r},0)"
    ws[f"L{r}"] = f"=IFERROR(H{r}/G{r},0)"
    ws[f"M{r}"] = f"=IFERROR(J{r}/F{r},0)"
    ws[f"N{r}"] = f"=IFERROR(J{r}/E{r},0)"
    ws[f"O{r}"] = f"=IFERROR(I{r}/J{r},0)"


def frame(ws, cols):
    """Headers, units, widths, number formats and forecast shading for a site or National tab."""
    last = list(cols)[-1]
    for col, (title, fmt, width) in cols.items():
        ws.column_dimensions[col].width = width
        c = header(ws, 13, ws[f"{col}1"].column, title)
        if col >= "R" or len(col) == 2:
            c.fill = PatternFill("solid", fgColor="7F7F7F")
        if col in UNITS:
            u = ws[f"{col}14"]
            u.value, u.font, u.alignment, u.fill = UNITS[col], Font(italic=True), CENTER, LIGHT
    ws.column_dimensions["Q"].width = 2
    ws.row_dimensions[13].height = 60
    spans = [("E", "Drivers", "E12:J12"), ("K", "Metrics", "K12:O12")]
    if last != "O":
        spans.append(("R", "Helper calculations (actuals, plan, same-month-last-year rates)", f"R12:{last}12"))
    for col, text, span in spans:
        ws.merge_cells(span)
        header(ws, 12, ws[f"{col}1"].column, text)
    ws.freeze_panes = f"E{FIRST_ROW}"
    ws.conditional_formatting.add(
        f"A{FIRST_ROW}:O{LAST_ROW}", FormulaRule(formula=[f'$D{FIRST_ROW}="Forecast"'], fill=LIGHT))
    for r in range(FIRST_ROW, LAST_ROW + 1):
        for col, (_, fmt, _) in cols.items():
            c = ws[f"{col}{r}"]
            if fmt:
                c.number_format = fmt
            if col >= "R" or len(col) == 2:
                c.font = Font(color="595959")
        if MONTHS[r - FIRST_ROW].month == 12:
            for col in "ABCDEFGHIJKLMNO":
                ws[f"{col}{r}"].border = Border(bottom=THIN)


def fy_block(ws):
    header(ws, 5, 1, "FY")
    header(ws, 5, 4, "Status")
    for col in "EFGHIJ" + METRIC_COLS:
        header(ws, 5, ws[f"{col}1"].column, SITE_COLS[col][0])
    ws.row_dimensions[5].height = 45
    for i, y in enumerate(YEARS):
        r = 6 + i
        ws[f"A{r}"] = y
        ws[f"A{r}"].font = BOLD
        ws[f"D{r}"] = (f'=IF(DATE(A{r},12,1)<=LastActual,"Actual",'
                       f'IF(DATE(A{r},1,1)>LastActual,"Forecast","Blended"))')
        for col in "EFGHIJ":
            ws[f"{col}{r}"] = f"=SUMIFS({RNG(col)},{RNG('B')},$A{r})"
        metric_formulas(ws, r)
        for col in "ADEFGHIJ" + METRIC_COLS:
            ws[f"{col}{r}"].fill = LIGHT
            if col in SITE_COLS and SITE_COLS[col][1]:
                ws[f"{col}{r}"].number_format = SITE_COLS[col][1]


def build_site(code, typ, state, host):
    ws = wb.create_sheet(code)
    ws["A1"] = f'=$B$2&"  |  {typ}  |  {state}' + (
        f'  |  payroll % of "&{CODE(host)}&" sales  |  volume from "&'
        + '&"+"&'.join(CODE(s) for s in SERVED[code]) + '&" sales"' if host else '"')
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "Code", "=" + CODE(code), "State", state
    ws["B2"].font = ws["D2"].font = BOLD
    ws["F2"] = "Leadership share of total payroll, last actual month"
    # from the raw actual helper columns, not column O: O's forecast rows depend on this cell
    lam = f"MATCH(LastActual,{RNG('A')},0)"
    ws["K2"] = f"=IFERROR(INDEX({RNG('V')},{lam})/(INDEX({RNG('U')},{lam})+INDEX({RNG('V')},{lam})),0)"
    ws["K2"].number_format = F_PCT
    ws["K2"].font = BOLD
    fy_block(ws)
    for i, m in enumerate(MONTHS):
        r = FIRST_ROW + i
        act = f'$D{r}="Actual"'
        fc = f'$D{r}="Forecast"'
        ly = r - 12 if i >= 12 else None  # same month last year
        ws[f"A{r}"] = m
        ws[f"B{r}"] = f"=YEAR(A{r})"
        ws[f"C{r}"] = f"=MONTH(A{r})"
        ws[f"D{r}"] = f'=IF(A{r}<=LastActual,"Actual","Forecast")'
        # drivers
        if typ == "IDC":
            ws[f"E{r}"] = f"='{host}'!E{r}"
            ws[f"R{r}"] = f'=IF({act},E{r},"")'
            ws[f"Y{r}"] = "=" + "+".join(f"'{s}'!E{r}" for s in SERVED[code])
        else:
            ws[f"E{r}"] = f"=IF({act},R{r},W{r})"
            ws[f"R{r}"] = f'=IF({act},{data_sum("C", r)},"")'
            ws[f"Y{r}"] = f"=E{r}"
        # forecast chain: cases / CPLH = hours -> x cost per hour = staff -> / (1 - leadership share) = total
        ws[f"F{r}"] = f"=IF({act},S{r},Cases_Forecast!{cf_col(3, SITE_IDX[code])}{CF_FIRST + i})"
        ws[f"G{r}"] = f"=IF({act},T{r},IFERROR(F{r}/K{r},0))"
        ws[f"H{r}"] = f"=IF({act},U{r},G{r}*L{r})"
        ws[f"J{r}"] = f"=IF({act},U{r}+V{r},IFERROR(H{r}/(1-AD{r}),0))"
        ws[f"I{r}"] = f"=IF({act},V{r},J{r}-H{r})"
        metric_formulas(ws, r)
        # CPLH and cost per hour are inputs in forecast months. Actual months read the raw actuals
        # (not F/G, H/G): G and H read K and L, and Excel flags that as circular even inside an IF
        ws[f"K{r}"] = f"=IF({act},IFERROR(S{r}/T{r},0),CPLH_Forecast!{cf_col(3, SITE_IDX[code])}{CF_FIRST + i})"
        ws[f"L{r}"] = f"=IF({act},IFERROR(U{r}/T{r},0),N(AE{r}))"
        # actuals and plan
        for col, src in (("S", "D"), ("T", "E"), ("U", "F"), ("V", "G")):
            ws[f"{col}{r}"] = f'=IF({act},{data_sum(src, r)},"")'
        ws[f"W{r}"] = f"=SUMIFS({PR('C')},{PR('A')},$A{r},{PR('B')},$B$2)"
        ws[f"Z{r}"] = (f"=INDEX(Calendar!$H$5:$L${4 + N},MATCH($A{r},Calendar!$A$5:$A${4 + N},0),"
                       f"MATCH($D$2,Calendar!$H$4:$L$4,0))")
        ws[f"AA{r}"] = f"=INDEX(Calendar!$M$5:$M${4 + N},MATCH($A{r},Calendar!$A$5:$A${4 + N},0))"
        # rates: from actuals, otherwise the same month last year (actual or forecast)
        for col, actual in (("AB", f"S{r}/Y{r}"), ("AC", f"U{r}/T{r}/Z{r}/AA{r}"), ("X", f"S{r}/T{r}")):
            ws[f"{col}{r}"] = f'=IF({act},IFERROR({actual},""),{f"{col}{ly}" if ly else chr(34) * 2})'
        ws[f"AD{r}"] = f'=IF({act},IFERROR(V{r}/(U{r}+V{r}),""),$K$2)'
        ws[f"AE{r}"] = f'=IFERROR(AC{r}*Z{r}*AA{r},"")'
        ws[f"AF{r}"] = f"=K{ly}" if ly else None
        ws[f"AG{r}"] = f'=IF({fc},IFERROR(K{r}/AF{r}-1,""),"")'
    frame(ws, SITE_COLS)


for s in SITES:
    build_site(*s)

# ---------------------------------------------------------------- National
ws = wb.create_sheet("National", index=wb.sheetnames.index(SITES[0][0]))
ws["A1"] = "National  |  8 regions + 2 IDCs  |  sales exclude IDCs"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "Model status"
ws["B2"] = "=Checks!B3"
ws["B2"].font = BOLD
fy_block(ws)
for i, m in enumerate(MONTHS):
    r = FIRST_ROW + i
    ws[f"A{r}"] = m
    ws[f"B{r}"] = f"=YEAR(A{r})"
    ws[f"C{r}"] = f"=MONTH(A{r})"
    ws[f"D{r}"] = f'=IF(A{r}<=LastActual,"Actual","Forecast")'
    ws[f"E{r}"] = "=" + "+".join(f"'{c}'!E{r}" for c in REGIONS)
    for col in "FGHIJ":
        ws[f"{col}{r}"] = "=" + "+".join(f"'{s[0]}'!{col}{r}" for s in SITES)
    metric_formulas(ws, r)
frame(ws, {k: v for k, v in SITE_COLS.items() if k <= "O" and len(k) == 1})

# site x FY summary tables to the right
ws.column_dimensions["Q"].width = 3
ws.column_dimensions["R"].width = 10
fy = lambda c, col, y: f"SUMIFS('{c}'!${col}${FIRST_ROW}:${col}${LAST_ROW},'{c}'!$B${FIRST_ROW}:$B${LAST_ROW},{y})"
tables = [
    ("Total payroll ($m)", lambda c, y: fy(c, "J", y), F_M),
    ("Cases picked (m)", lambda c, y: fy(c, "F", y), F_CASES),
    ("Payroll % of sales (IDC: % of host sales)", lambda c, y: f"IFERROR({fy(c, 'J', y)}/{fy(c, 'E', y)},0)", F_PCT),
    ("CPLH (cases / hr)", lambda c, y: f"IFERROR({fy(c, 'F', y)}/{fy(c, 'G', y)},0)", F_1),
    ("Cost per case ($, total payroll)", lambda c, y: f"IFERROR({fy(c, 'J', y)}/{fy(c, 'F', y)},0)", F_CPC),
]
row = 4
for title, fn, fmt in tables:
    ws.cell(row=row, column=18, value=title).font = BOLD
    header(ws, row + 1, 18, "Site")
    for k, y in enumerate(YEARS):
        header(ws, row + 1, 19 + k, f"FY{y}", width=11)
    for i, (code, *_rest) in enumerate(SITES):
        ws.cell(row=row + 2 + i, column=18, value="=" + CODE(code))
        for k, y in enumerate(YEARS):
            ws.cell(row=row + 2 + i, column=19 + k, value="=" + fn(code, y)).number_format = fmt
    row += len(SITES) + 4

# ------------------------------------------------------------------ Checks
ws = wb.create_sheet("Checks", index=wb.sheetnames.index("National"))
ws.column_dimensions["A"].width = 62
for col in "BCD":
    ws.column_dimensions[col].width = 16
ws["A1"] = "Checks"
ws["A1"].font = Font(bold=True, size=14)
ws["A3"] = "Overall"
ws["B3"] = '=IF(COUNTIF(D6:D80,"ERROR")=0,"ALL CHECKS OK","CHECK ERRORS")'
ws["B3"].font = BOLD
ws["C3"] = '=COUNTIF(D6:D80,"WARN")&" warning(s)"'
for j, t in enumerate(["Check", "Value", "Expected", "Result"], 1):
    header(ws, 5, j, t)
checks = []
act_months = '(YEAR(LastActual)-2025)*12+MONTH(LastActual)'
fc_months = '(2028-YEAR(LastActual))*12+12-MONTH(LastActual)'
for code, *_r in SITES:
    checks.append((f'={CODE(code)}&": actual rows in Data_Actuals (one per month)"',
                   f'=COUNTIFS({DR("B")},{CODE(code)},{DR("A")},"<="&LastActual)', f"={act_months}", "eq"))
for code in REGIONS:
    checks.append((f'={CODE(code)}&": plan sales present for every forecast month"',
                   f'=COUNTIFS({PR("B")},{CODE(code)},{PR("A")},">"&LastActual,{PR("C")},">0")',
                   f"={fc_months}", "eq"))
checks.append(("Unknown site codes in Data_Actuals",
               f"=SUMPRODUCT(({DR('B')}<>\"\")*(COUNTIF(SiteCodes,{DR('B')})=0))", "=0", "eq"))
checks.append(("At least 12 actual months (same month last year needs them)", f"={act_months}", "=12", "ge"))
checks.append(("National sales = sum of 8 regions (IDCs excluded)",
               f"=SUM(National!E{FIRST_ROW}:E{LAST_ROW})",
               "=" + "+".join(f"SUM('{c}'!E{FIRST_ROW}:E{LAST_ROW})" for c in REGIONS), "eq"))
checks.append(("National total payroll = sum of all 10 sites",
               f"=SUM(National!J{FIRST_ROW}:J{LAST_ROW})",
               "=" + "+".join(f"SUM('{s[0]}'!J{FIRST_ROW}:J{LAST_ROW})" for s in SITES), "eq"))
checks.append(("Staff + leadership = total payroll (national, all months)",
               f"=SUM(National!H{FIRST_ROW}:I{LAST_ROW})", f"=SUM(National!J{FIRST_ROW}:J{LAST_ROW})", "eq"))
for tab in ("Cases_Forecast", "CPLH_Forecast"):
    for b, what in ((1, "adjustments"), (2, "overrides")):
        checks.append((f"{tab}: {what} typed in actual months (ignored)",
                       f'=SUMPRODUCT(({CF_STATUS(tab)}="Actual")*({CF_RANGE(tab, b)}<>""))', "=0", "warn"))
for code, *_r in SITES:
    rng = f"'{code}'!AG{FIRST_ROW}:AG{LAST_ROW}"
    checks.append((f'={CODE(code)}&": forecast months where CPLH moves more than the tolerance vs LY"',
                   f'=COUNTIF({rng},">"&CPLHTol)+COUNTIF({rng},"<"&-CPLHTol)', "=0", "warn"))
for i, (label, val, exp, op) in enumerate(checks):
    r = 6 + i
    ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = label, val, exp
    # warn checks are counts, so compare exactly; eq checks allow $ rounding
    test = {"ge": f"B{r}>=C{r}", "warn": f"B{r}=C{r}"}.get(op, f"ABS(B{r}-C{r})<=MAX(1,ABS(C{r})*0.000001)")
    ws[f"D{r}"] = f'=IF({test},"OK","{"WARN" if op == "warn" else "ERROR"}")'
    for col in "BC":
        ws[f"{col}{r}"].number_format = F_INT
ws.conditional_formatting.add("D6:D80", FormulaRule(formula=['D6="ERROR"'], fill=PatternFill("solid", fgColor="F8CBAD")))
ws.conditional_formatting.add("D6:D80", FormulaRule(formula=['D6="OK"'], fill=PatternFill("solid", fgColor="C6EFCE")))
ws.conditional_formatting.add("D6:D80", FormulaRule(formula=['D6="WARN"'], fill=PatternFill("solid", fgColor="FFEB9C")))

# tab colours
for nm in wb.sheetnames:
    if nm in ("Inputs", "Data_Actuals", "Plan_Inputs", "Holidays", "Cases_Forecast", "CPLH_Forecast", "CLR"):
        wb[nm].sheet_properties.tabColor = "FFC000"
    elif nm in REGIONS:
        wb[nm].sheet_properties.tabColor = "1F3864"
    elif nm in SERVED:
        wb[nm].sheet_properties.tabColor = "2E75B6"
    elif nm == "National":
        wb[nm].sheet_properties.tabColor = "00B050"

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
