"""Build the warehouse payroll forecast model (synthetic data) as an Excel workbook.

Run:  .venv/bin/python build_model.py
Output: warehouse_payroll_forecast.xlsx
"""
import calendar
import math
import random
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "warehouse_payroll_forecast.xlsx"
# bounded ranges (whole-column refs are slow); room for pasted data up to these rows
DATA_MAX, HOL_MAX = 5000, 1000
DR = lambda col: f"Data_Actuals!${col}$2:${col}${DATA_MAX}"
PR = lambda col: f"Plan_Inputs!${col}$2:${col}${DATA_MAX}"
HR = lambda col: f"Holidays!${col}$2:${col}${HOL_MAX}"
START = date(2025, 1, 1)
END = date(2028, 12, 1)
LAST_ACTUAL = date(2026, 9, 1)
MONTHS = []
d = START
while d <= END:
    MONTHS.append(d)
    d = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
N = len(MONTHS)  # 48
FIRST_ROW = 15
LAST_ROW = FIRST_ROW + N - 1  # 62
YEARS = [2025, 2026, 2027, 2028]

# code, type, state, host region (IDCs borrow host sales)
SITES = [
    ("NSW-1", "Region", "NSW", None),
    ("VIC-1", "Region", "VIC", None),
    ("QLD-1", "Region", "QLD", None),
    ("NSW-2", "Region", "NSW", None),
    ("VIC-2", "Region", "VIC", None),
    ("QLD-2", "Region", "QLD", None),
    ("SA-1", "Region", "SA", None),
    ("WA-1", "Region", "WA", None),
    ("IDC-NSW", "IDC", "NSW", "NSW-1"),
    ("IDC-VIC", "IDC", "VIC", "VIC-1"),
]
REGIONS = [s[0] for s in SITES if s[1] == "Region"]
# regions each IDC ships to: their sales drive IDC volume (payroll % stays on host sales)
SERVED = {"IDC-NSW": ["NSW-1", "NSW-2"], "IDC-VIC": ["VIC-1", "VIC-2"]}
STATES = ["NSW", "VIC", "QLD", "SA", "WA"]
# every site code in the workbook reads from the Code column of the Sites table on Inputs,
# so typing new codes there renames the sites everywhere (tab names are cosmetic)
SITE_ROW = {s[0]: 36 + i for i, s in enumerate(SITES)}
CODE = lambda c: f"Inputs!$A${SITE_ROW[c]}"

WAGE_TABLE = [(date(2026, 1, 1), 0.040), (date(2027, 1, 1), 0.045), (date(2028, 1, 1), 0.045)]
PH_ROSTER = 0.5   # share of a normal day's hours worked on a public holiday
PH_LOADING = 1.5  # extra pay on a public holiday (1.5 = paid 250%)


# ---------------------------------------------------------------- holidays
def easter(y):
    a, b, c = y % 19, y // 100, y % 100
    dd, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - dd - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    month = (h + l - 7 * m + 114) // 31
    day = (h + l - 7 * m + 114) % 31 + 1
    return date(y, month, day)


def nth_weekday(y, m, wd, n):
    """n-th weekday (0=Mon) of a month; n=-1 for the last."""
    if n > 0:
        d0 = date(y, m, 1)
        return d0 + timedelta(days=(wd - d0.weekday()) % 7 + 7 * (n - 1))
    d0 = date(y, m, calendar.monthrange(y, m)[1])
    return d0 - timedelta(days=(d0.weekday() - wd) % 7)


def next_monday(d0):
    return d0 + timedelta(days=(7 - d0.weekday()) % 7)


def holidays():
    rows = []
    for y in YEARS:
        ny = date(y, 1, 1)
        aus = date(y, 1, 26)
        gf = easter(y) - timedelta(days=2)
        xmas = date(y, 12, 25)
        anzac = date(y, 4, 25)
        for st in STATES:
            h = [(ny, "New Year's Day")]
            if ny.weekday() >= 5:
                h.append((next_monday(ny), "New Year's Day (additional)"))
            h.append((next_monday(aus) if aus.weekday() >= 5 else aus, "Australia Day"))
            h.append((gf, "Good Friday"))
            if st != "WA":
                h.append((gf + timedelta(days=1), "Easter Saturday"))
            h.append((gf + timedelta(days=2), "Easter Sunday"))
            h.append((gf + timedelta(days=3), "Easter Monday"))
            h.append((anzac, "Anzac Day"))
            if (st == "WA" and anzac.weekday() >= 5) or (st == "SA" and anzac.weekday() == 6):
                h.append((next_monday(anzac), "Anzac Day (additional)"))
            if st in ("NSW", "VIC", "SA"):
                h.append((nth_weekday(y, 6, 0, 2), "King's Birthday"))
            elif st == "QLD":
                h.append((nth_weekday(y, 10, 0, 1), "King's Birthday"))
            else:
                h.append((nth_weekday(y, 9, 0, -1), "King's Birthday"))
            labour = {"NSW": (10, 1), "SA": (10, 1), "VIC": (3, 2), "QLD": (5, 1), "WA": (3, 1)}[st]
            h.append((nth_weekday(y, labour[0], 0, labour[1]), "Labour Day"))
            if st == "VIC":
                h.append((nth_weekday(y, 11, 1, 1), "Melbourne Cup"))
                h.append((nth_weekday(y, 9, 5, -1) - timedelta(days=1), "AFL Grand Final Friday (check date)"))
            if st == "SA":
                h.append((nth_weekday(y, 3, 0, 2), "Adelaide Cup"))
            if st == "WA":
                h.append((nth_weekday(y, 6, 0, 1), "WA Day"))
            h.append((xmas, "Christmas Day"))
            h.append((date(y, 12, 26), "Proclamation Day" if st == "SA" else "Boxing Day"))
            if xmas.weekday() == 5:
                h += [(date(y, 12, 27), "Christmas (additional)"), (date(y, 12, 28), "Boxing Day (additional)")]
            elif xmas.weekday() == 6:
                h.append((date(y, 12, 27), "Christmas (additional)"))
            elif xmas.weekday() == 4:
                h.append((date(y, 12, 28), "Boxing Day (additional)"))
            rows += [(dt, name, st) for dt, name in h]
    rows.sort(key=lambda r: (r[2], r[0]))
    return rows


HOLIDAYS = holidays()


def ph_factor(month, state):
    days = calendar.monthrange(month.year, month.month)[1]
    h = sum(1 for dt, _, st in HOLIDAYS if st == state and dt.year == month.year and dt.month == month.month)
    return (days - h + h * PH_ROSTER * (1 + PH_LOADING)) / (days - h + h * PH_ROSTER)


def wage_index(month):
    idx = 1.0
    for eff, pct in WAGE_TABLE:
        if eff <= month:
            idx *= 1 + pct
    return idx


# ---------------------------------------------------------- synthetic data
SEASON_SALES = {1: .97, 2: .90, 3: 1.0, 4: .99, 5: .98, 6: .96, 7: .98, 8: .97, 9: .96, 10: 1.0, 11: 1.03, 12: 1.25}
SEASON_CPLH = {1: .97, 2: .99, 3: 1.0, 4: .99, 5: 1.0, 6: 1.0, 7: 1.0, 8: 1.0, 9: 1.0, 10: 1.01, 11: 1.03, 12: 1.06}
SEASON_LEAD = {11: 1.05, 12: 1.15}
STATE_CPH = {"NSW": 66.0, "VIC": 65.0, "QLD": 63.0, "SA": 61.0, "WA": 69.0}
# monthly sales $m (2025), cases per $ sales, CPLH, leadership $m per month
PARAMS = {
    "NSW-1": (150, .0210, 78, .32), "VIC-1": (140, .0215, 80, .30), "QLD-1": (115, .0220, 74, .26),
    "NSW-2": (125, .0212, 76, .28), "VIC-2": (135, .0208, 79, .30), "QLD-2": (120, .0218, 75, .27),
    "SA-1": (80, .0225, 72, .21), "WA-1": (105, .0223, 73, .25),
    "IDC-NSW": (None, .0030, 110, .15), "IDC-VIC": (None, .0028, 105, .14),
}


def synth():
    rng = random.Random(42)
    data = {}
    for code, typ, state, host in SITES:
        base_sales, cps0, cplh0, lead0 = PARAMS[code]
        for t, m in enumerate(MONTHS):
            if typ == "Region":
                sales = base_sales * 1e6 * SEASON_SALES[m.month] * 1.035 ** (t / 12) * (1 + rng.gauss(0, .015))
            else:
                sales = data[(host, m)]["sales"]
            cps = cps0 * 0.985 ** (t / 12) * (0.98 if m.month == 12 else 1) * (1 + rng.gauss(0, .01))
            driver = sum(data[(s, m)]["sales"] for s in SERVED[code]) if typ == "IDC" else sales
            cases = driver * cps
            cplh = cplh0 * SEASON_CPLH[m.month] * 1.01 ** (t / 12) * (1 + rng.gauss(0, .015))
            hours = cases / cplh
            cph = STATE_CPH[state] * (1.06 if typ == "IDC" else 1) * wage_index(m) * ph_factor(m, state) * (1 + rng.gauss(0, .008))
            staff = hours * cph
            lead = lead0 * 1e6 * wage_index(m) * SEASON_LEAD.get(m.month, 1) * (1 + rng.gauss(0, .02))
            data[(code, m)] = dict(sales=sales, cases=cases, hours=hours, staff=staff, lead=lead)
    return data


DATA = synth()

# ------------------------------------------------------------------ styles
NAVY = PatternFill("solid", fgColor="1F3864")
LIGHT = PatternFill("solid", fgColor="DDEBF7")
INPUT = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="F2F2F2")
WHITE_B = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
BLUE = Font(color="0000FF")
HELP = Font(color="808080", italic=True)
THIN = Side(style="thin", color="1F3864")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

F_M = '#,##0.0,,'        # $ or units shown in millions
F_CASES = '#,##0.00,,'
F_INT = '#,##0'
F_1 = '0.0'
F_CPH = '$#,##0.00'
F_CPC = '$#,##0.000'
F_PCT = '0.00%'
F_DATE = 'mmm-yy'
F_RATIO = '0.00000'

wb = Workbook()


def header(ws, row, col, text, fill=NAVY, font=WHITE_B, width=None):
    c = ws.cell(row=row, column=col, value=text)
    c.fill, c.font, c.alignment = fill, font, CENTER
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return c


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


# ------------------------------------------------------------------ README
ws = wb.active
ws.title = "README"
readme = [
    ("Warehouse payroll forecast model", BOLD),
    ("Synthetic data only. Monthly, Jan 2025 to Dec 2028.", None),
    ("", None),
    ("HOW TO UPDATE ACTUALS", BOLD),
    ("1. Paste the new actuals over the rows in Data_Actuals (keep the header and column order). One row per site per month.", None),
    ("   Months must be the 1st of the month. IDC rows leave Sales blank.", None),
    ("2. Set 'Last actual month' on the Inputs tab. Every tab switches actual/forecast from that date.", None),
    ("3. Paste the plan (sales forecast + payroll % of sales) into Plan_Inputs, same layout.", None),
    ("4. Check the Checks tab says ALL CHECKS OK.", None),
    ("", None),
    ("FORECAST LOGIC (per site, per month)", BOLD),
    ("Sales = plan sales x (1 + scenario sales lever). IDCs use their host region's sales (Sites table on Inputs).", None),
    ("Total payroll = sales x plan payroll % (the driver).", None),
    ("Leadership = sales x leadership % run rate, or run-rate $ stepped up by wage increases (switch on Inputs).", None),
    ("Staff payroll = total payroll - leadership.", None),
    ("Cost per hour = run rate x public-holiday factor x wage index.  Hours = staff payroll / cost per hour.", None),
    ("Cases = volume driver sales x cases-per-$ run rate x (1 + scenario cases lever).", None),
    ("Volume driver: a region's own sales; an IDC's served regions (Sites table on Inputs).", None),
    ("Each run rate can drift by a yearly trend (Inputs), e.g. price inflation lowering cases per $ of sales.", None),
    ("CPLH = cases / hours.  Cost per hour = staff payroll / hours (excludes leadership). Cost per case = total payroll / cases (includes leadership).", None),
    ("", None),
    ("BOTTOM-UP CROSS-CHECK", BOLD),
    ("Bottom-up staff payroll = cases / (LY CPLH x (1 + productivity target)) x forecast cost per hour.", None),
    ("If top-down (plan) payroll is well below bottom-up, the plan needs a productivity jump to hit. See National tab", None),
    ("and the Checks tab, which warns when forecast CPLH moves more than the tolerance vs the same month last year.", None),
    ("", None),
    ("RUN RATES (cases per $ sales, cost per hour, leadership %)", BOLD),
    ("Seasonal index = each calendar month vs the average over the last 12 actual months.", None),
    ("Level = average of the last 6 actual months / average seasonal index of those months.", None),
    ("Forecast = level x seasonal index, so the last 6 months set the level and last year sets the shape.", None),
    ("Cost per hour is first stripped of public-holiday and wage effects, then they are added back for each forecast month.", None),
    ("So Easter moving from April to March, or an extra holiday, changes cost per hour in that month.", None),
    ("", None),
    ("NATIONAL", BOLD),
    ("Sales = 8 regions only (IDC sales are a copy of the host's and are excluded). Cases, hours and payroll = all 10 sites.", None),
    ("", None),
    ("SITE CODES", BOLD),
    ("Type your own codes in the Code column of the Sites table on Inputs; every tab follows. Renaming the tabs is optional.", None),
    ("", None),
    ("COLOURS", BOLD),
    ("Yellow cells with blue text are inputs. Blue-shaded rows are forecast months. Grey columns are helper calculations.", None),
    ("Public holidays are generated by rule; check against the official state lists (AFL Grand Final Friday in particular).", None),
]
for i, (txt, font) in enumerate(readme, 1):
    c = ws.cell(row=i, column=1, value=txt)
    if font:
        c.font = font
ws.column_dimensions["A"].width = 120

# ------------------------------------------------------------------ Inputs
ws = wb.create_sheet("Inputs")
ws.column_dimensions["A"].width = 34
for col in "BCDEF":
    ws.column_dimensions[col].width = 14
ws["A1"] = "Inputs"
ws["A1"].font = Font(bold=True, size=14)

ws["A3"] = "Last actual month"
ws["B3"] = LAST_ACTUAL
ws["B3"].number_format = F_DATE
ws["A4"] = "Run-rate months (level)"
ws["B4"] = 6
ws["A5"] = "Seasonality window (months)"
ws["B5"] = 12
ws["A6"] = "Model status"
ws["B6"] = "=Checks!B3"
ws["B6"].font = BOLD
for r in (3, 4, 5):
    ws.cell(row=r, column=2).fill, ws.cell(row=r, column=2).font = INPUT, BLUE
name("LastActual", "Inputs!$B$3")
name("RunRateMonths", "Inputs!$B$4")
name("SeasonMonths", "Inputs!$B$5")

ws.column_dimensions["D"].width = 40
ws["D3"] = "Run-rate trend per year (from last actual month)"
ws["D3"].font = BOLD
for i, (label, nm, v) in enumerate([("Cases per $ sales (price inflation pushes this down)", "TrendCPS", -0.015),
                                    ("Cost per hour (before wage increases)", "TrendCPH", 0.0),
                                    ("Leadership", "TrendLead", 0.0)]):
    r = 4 + i
    ws[f"D{r}"], ws[f"E{r}"] = label, v
    ws[f"E{r}"].number_format, ws[f"E{r}"].fill, ws[f"E{r}"].font = '0.0%', INPUT, BLUE
    name(nm, f"Inputs!$E${r}")

ws["D12"] = "Leadership forecast method"
ws["D12"].font = BOLD
ws["E12"] = "% of sales"
ws["E12"].fill, ws["E12"].font = INPUT, BLUE
lead_dv = DataValidation(type="list", formula1='"% of sales,$ + wage increases"', allow_blank=False)
ws.add_data_validation(lead_dv)
lead_dv.add("E12")
ws["D13"] = "% of sales: grows with sales. $ + wage increases: run-rate dollars that step up with each wage increase."
ws["D13"].font = HELP
name("LeadMethod", "Inputs!$E$12")

ws["D8"] = "Bottom-up cross-check"
ws["D8"].font = BOLD
for i, (label, nm, v) in enumerate([("Productivity target: CPLH change vs same month LY", "ProdTarget", 0.0),
                                    ("Warn if forecast CPLH moves more than this vs LY", "CPLHTol", 0.10)]):
    r = 9 + i
    ws[f"D{r}"], ws[f"E{r}"] = label, v
    ws[f"E{r}"].number_format, ws[f"E{r}"].fill, ws[f"E{r}"].font = '0.0%', INPUT, BLUE
    name(nm, f"Inputs!$E${r}")

ws["A8"] = "Scenario"
ws["A8"].font = BOLD
ws["A9"] = "Selected scenario"
ws["B9"] = "Base"
ws["B9"].fill, ws["B9"].font = INPUT, BLUE
dv = DataValidation(type="list", formula1='"Base,Upside,Downside"', allow_blank=False)
ws.add_data_validation(dv)
dv.add("B9")
for c, t in zip("ABC", ["Scenario", "Sales lever", "Cases lever"]):
    header(ws, 11, " ABC".index(c), t)
for i, (sc, s, cs) in enumerate([("Base", 0, 0), ("Upside", .03, .02), ("Downside", -.03, -.02)]):
    r = 12 + i
    ws.cell(row=r, column=1, value=sc)
    for col, v in ((2, s), (3, cs)):
        c = ws.cell(row=r, column=col, value=v)
        c.number_format, c.fill, c.font = '0.0%', INPUT, BLUE
ws["A16"] = "Sales lever (active)"
ws["B16"] = "=INDEX($B$12:$B$14,MATCH($B$9,$A$12:$A$14,0))"
ws["A17"] = "Cases lever (active)"
ws["B17"] = "=INDEX($C$12:$C$14,MATCH($B$9,$A$12:$A$14,0))"
ws["B16"].number_format = ws["B17"].number_format = '0.0%'
name("SalesAdj", "Inputs!$B$16")
name("CasesAdj", "Inputs!$B$17")

ws["A19"] = "Wage increases (applied to cost per hour)"
ws["A19"].font = BOLD
header(ws, 20, 1, "Effective month")
header(ws, 20, 2, "Increase")
for i in range(8):
    r = 21 + i
    a, b = ws.cell(row=r, column=1), ws.cell(row=r, column=2)
    if i < len(WAGE_TABLE):
        a.value, b.value = WAGE_TABLE[i]
    a.number_format, b.number_format = F_DATE, '0.0%'
    a.fill = b.fill = INPUT
    a.font = b.font = BLUE
ws["C21"] = "History row: lets the model strip past increases out of actual cost per hour. Replace with your real ones."
ws["C21"].font = HELP
name("WageDates", "Inputs!$A$21:$A$28")
name("WageRates", "Inputs!$B$21:$B$28")

ws["A30"] = "Public holidays"
ws["A30"].font = BOLD
ws["A31"] = "Share of normal hours worked on a PH"
ws["B31"] = PH_ROSTER
ws["A32"] = "PH pay loading (1.5 = paid 250%)"
ws["B32"] = PH_LOADING
for r in (31, 32):
    ws.cell(row=r, column=2).fill, ws.cell(row=r, column=2).font = INPUT, BLUE
name("PH_Roster", "Inputs!$B$31")
name("PH_Loading", "Inputs!$B$32")

ws["A34"] = "Sites"
ws["A34"].font = BOLD
ws["B34"] = "Type your own codes in the Code column; every tab follows. Renaming the tabs to match is optional."
ws["B34"].font = HELP
for j, t in enumerate(["Code", "Type", "State", "Payroll % basis", "In national sales?", "Volume driver"], 1):
    header(ws, 35, j, t)
for i, (code, typ, st, host) in enumerate(SITES):
    r = 36 + i
    vals = [code, typ, st, f"=A{SITE_ROW[host or code]}", "No" if typ == "IDC" else "Yes",
            "=" + '&"+"&'.join(f"A{SITE_ROW[s]}" for s in SERVED.get(code, [code]))]
    for j, v in enumerate(vals, 1):
        ws.cell(row=r, column=j, value=v)
    ws.cell(row=r, column=1).fill, ws.cell(row=r, column=1).font = INPUT, BLUE
name("SiteCodes", f"Inputs!$A$36:$A${35 + len(SITES)}")

# ------------------------------------------------------------ Data_Actuals
ws = wb.create_sheet("Data_Actuals")
cols = ["Month", "Site", "Sales $", "Cases picked", "WH hours worked", "Staff payroll $", "Leadership payroll $"]
for j, t in enumerate(cols, 1):
    header(ws, 1, j, t, width=16)
r = 2
for m in MONTHS:
    if m > LAST_ACTUAL:
        break
    for code, typ, _, _ in SITES:
        v = DATA[(code, m)]
        row = [m, "=" + CODE(code), None if typ == "IDC" else round(v["sales"]), round(v["cases"]), round(v["hours"]),
               round(v["staff"]), round(v["lead"])]
        for j, val in enumerate(row, 1):
            c = ws.cell(row=r, column=j, value=val)
            c.number_format = F_DATE if j == 1 else F_INT
        r += 1
ws.freeze_panes = "A2"

# ------------------------------------------------------------- Plan_Inputs
ws = wb.create_sheet("Plan_Inputs")
for j, t in enumerate(["Month", "Site", "Sales forecast $", "Payroll % of sales (total)"], 1):
    header(ws, 1, j, t, width=22)
ws["F1"] = "Stand-in for the long-range plan. IDC payroll % is a % of the host region's sales (see Sites on Inputs)."
ws["F1"].font = HELP
r = 2
for m in MONTHS:
    if m.year < 2026:
        continue
    for code, typ, _, host in SITES:
        v = DATA[(code, m)]
        pct = round((v["staff"] + v["lead"]) / v["sales"], 4)
        row = [m, "=" + CODE(code), None if typ == "IDC" else round(v["sales"], -3), pct]
        for j, val in enumerate(row, 1):
            c = ws.cell(row=r, column=j, value=val)
            c.number_format = [F_DATE, None, F_INT, F_PCT][j - 1] or "General"
        r += 1
ws.freeze_panes = "A2"

# ---------------------------------------------------------------- Holidays
ws = wb.create_sheet("Holidays")
for j, t in enumerate(["Date", "Holiday", "State"], 1):
    header(ws, 1, j, t, width=[12, 38, 8][j - 1])
for i, (dt, nm, st) in enumerate(HOLIDAYS, 2):
    ws.cell(row=i, column=1, value=dt).number_format = "dd-mmm-yy"
    ws.cell(row=i, column=2, value=nm)
    ws.cell(row=i, column=3, value=st)
ws["E1"] = "Generated by rule. Add, remove or move rows freely; the Calendar tab counts them by state and month."
ws["E1"].font = HELP
ws.freeze_panes = "A2"

# ---------------------------------------------------------------- Calendar
ws = wb.create_sheet("Calendar")
ws["A1"] = "Calendar: public holidays, PH cost factor, wage index"
ws["A1"].font = BOLD
ws["A2"] = "PH factor = average cost per hour in the month vs a month with no public holidays."
ws["A2"].font = HELP
heads = ["Month", "Days"] + [f"PH {s}" for s in STATES] + STATES + ["Wage index"]
for j, t in enumerate(heads, 1):
    header(ws, 4, j, t, width=11)
ws.cell(row=3, column=3, value="Public holidays (count)").font = BOLD
ws.cell(row=3, column=8, value="PH cost factor").font = BOLD
for i, m in enumerate(MONTHS):
    r = 5 + i
    ws.cell(row=r, column=1, value=m).number_format = F_DATE
    ws.cell(row=r, column=2, value=f"=DAY(EOMONTH(A{r},0))")
    for k, st in enumerate(STATES):
        cnt = 3 + k
        ws.cell(row=r, column=cnt, value=(
            f'=COUNTIFS({HR("A")},">="&$A{r},{HR("A")},"<="&EOMONTH($A{r},0),{HR("C")},"{st}")'))
        h = get_column_letter(cnt)
        c = ws.cell(row=r, column=8 + k, value=(
            f"=($B{r}-{h}{r}+{h}{r}*PH_Roster*(1+PH_Loading))/($B{r}-{h}{r}+{h}{r}*PH_Roster)"))
        c.number_format = '0.0000'
    # product of (1 + increase) for every increase effective on or before this month
    c = ws.cell(row=r, column=13, value=f"=EXP(SUMPRODUCT((WageDates<=$A{r})*(WageDates<>\"\")*LN(1+WageRates)))")
    c.number_format = '0.0000'
ws.freeze_panes = "B5"

# -------------------------------------------------------------- site tabs
SITE_COLS = {
    "A": ("Month", F_DATE, 9), "B": ("Year", "0", 6), "C": ("Mth", "0", 5), "D": ("Status", None, 9),
    "E": ("Sales", F_M, 10), "F": ("Cases picked", F_CASES, 10), "G": ("WH hours worked", F_INT, 11),
    "H": ("Staff payroll", F_M, 10), "I": ("Leadership payroll", F_M, 11), "J": ("Total payroll", F_M, 10),
    "K": ("CPLH (cases / hr)", F_1, 10), "L": ("Cost per hour", F_CPH, 10), "M": ("Cost per case", F_CPC, 10),
    "N": ("Staff payroll % sales", F_PCT, 10), "O": ("Leadership % sales", F_PCT, 10), "P": ("Total payroll % sales", F_PCT, 10),
    "R": ("Act sales", F_INT, 11), "S": ("Act cases", F_INT, 11), "T": ("Act hours", F_INT, 10),
    "U": ("Act staff $", F_INT, 11), "V": ("Act leadership $", F_INT, 11),
    "W": ("Act cases per $", F_RATIO, 10), "X": ("Act CPH (ex PH & wage)", F_CPH, 11), "Y": ("Act leadership %", F_PCT, 10),
    "Z": ("PH factor", '0.0000', 8), "AA": ("Wage index", '0.0000', 8),
    "AB": ("Plan sales", F_INT, 12), "AC": ("Plan payroll %", F_PCT, 9),
    "AD": ("Season idx cases/$", '0.000', 9), "AE": ("Season idx CPH", '0.000', 9), "AF": ("Season idx lead %", '0.000', 9),
    "AG": ("Fcst cases per $", F_RATIO, 10), "AH": ("Fcst CPH", F_CPH, 9), "AI": ("Fcst leadership %", F_PCT, 9),
    "AJ": ("Years after last actual", '0.00', 8),
    "AK": ("LY CPLH (same month)", F_1, 9), "AL": ("Bottom-up hours", F_INT, 11),
    "AM": ("Bottom-up staff payroll", F_INT, 12), "AN": ("Top-down minus bottom-up $", F_INT, 12),
    "AO": ("Fcst CPLH vs LY", F_PCT, 9),
    "AP": ("Volume driver sales", F_INT, 12),
    "AQ": ("Act leadership $ (ex wage)", F_INT, 12), "AR": ("Season idx lead $", '0.000', 9),
    "AS": ("Fcst leadership $", F_INT, 12),
}
UNITS = {"E": "$m", "F": "m", "G": "#", "H": "$m", "I": "$m", "J": "$m", "K": "#", "L": "$", "M": "$",
         "N": "%", "O": "%", "P": "%"}
RNG = lambda col: f"${col}${FIRST_ROW}:${col}${LAST_ROW}"
ACT_WIN = lambda months: f'{RNG("A")},">"&EDATE(LastActual,-{months}),{RNG("A")},"<="&LastActual'


def data_sum(col, r):
    return f"SUMIFS({DR(col)},{DR('A')},$A{r},{DR('B')},$B$2)"


def style_site_sheet(ws):
    for col, (title, fmt, width) in SITE_COLS.items():
        ws.column_dimensions[col].width = width
        c = header(ws, 13, ws[f"{col}1"].column, title)
        if col >= "R" or len(col) == 2:
            c.fill = PatternFill("solid", fgColor="7F7F7F")
        if col in UNITS:
            u = ws[f"{col}14"]
            u.value, u.font, u.alignment, u.fill = UNITS[col], Font(italic=True), CENTER, LIGHT
    ws.column_dimensions["Q"].width = 2
    ws.row_dimensions[13].height = 45
    for col, text, span in (("E", "Drivers", "E12:J12"), ("K", "Metrics", "K12:P12"),
                            ("R", "Helper calculations (actuals, run rates, plan)", f"R12:{list(SITE_COLS)[-1]}12")):
        ws.merge_cells(span)
        header(ws, 12, ws[f"{col}1"].column, text)
    ws.freeze_panes = f"E{FIRST_ROW}"
    ws.conditional_formatting.add(
        f"A{FIRST_ROW}:P{LAST_ROW}", FormulaRule(formula=[f'$D{FIRST_ROW}="Forecast"'], fill=LIGHT))
    for r in range(FIRST_ROW, LAST_ROW + 1):
        for col, (_, fmt, _) in SITE_COLS.items():
            c = ws[f"{col}{r}"]
            if fmt:
                c.number_format = fmt
            if col >= "R" or len(col) == 2:
                c.font = Font(color="595959")
        if MONTHS[r - FIRST_ROW].month == 12:
            for col in "ABCDEFGHIJKLMNOP":
                ws[f"{col}{r}"].border = Border(bottom=THIN)


def metric_formulas(ws, r):
    ws[f"K{r}"] = f"=IFERROR(F{r}/G{r},0)"
    ws[f"L{r}"] = f"=IFERROR(H{r}/G{r},0)"
    ws[f"M{r}"] = f"=IFERROR(J{r}/F{r},0)"
    ws[f"N{r}"] = f"=IFERROR(H{r}/E{r},0)"
    ws[f"O{r}"] = f"=IFERROR(I{r}/E{r},0)"
    ws[f"P{r}"] = f"=IFERROR(J{r}/E{r},0)"


def fy_block(ws):
    header(ws, 5, 1, "FY")
    header(ws, 5, 4, "Status")
    for col in "EFGHIJKLMNOP":
        header(ws, 5, ws[f"{col}1"].column, SITE_COLS[col][0])
    ws.row_dimensions[5].height = 45
    for i, y in enumerate(YEARS):
        r = 6 + i
        ws[f"A{r}"] = y
        ws[f"A{r}"].font = BOLD
        ws[f"D{r}"] = (f'=IF(DATE(A{r},12,1)<=LastActual,"Actual",'
                       f'IF(DATE(A{r},1,1)>LastActual,"Forecast","Blended"))')
        for col in "EFGHIJ":
            ws[f"{col}{r}"] = f"=SUMIFS({RNG(col)},{RNG('B')},$A{r})"
            ws[f"{col}{r}"].number_format = SITE_COLS[col][1]
        metric_formulas(ws, r)
        for col in "KLMNOP":
            ws[f"{col}{r}"].number_format = SITE_COLS[col][1]
        for col in "ADEFGHIJKLMNOP":
            ws[f"{col}{r}"].fill = LIGHT


def build_site(code, typ, state, host):
    ws = wb.create_sheet(code)
    ws["A1"] = f'=$B$2&"  |  {typ}  |  {state}' + (
        f'  |  payroll % of "&{CODE(host)}&" sales  |  volume from "&'
        + '&"+"&'.join(CODE(s) for s in SERVED[code]) + '&" sales"' if host else '"')
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "Code", "=" + CODE(code), "State", state
    ws["B2"].font = ws["D2"].font = BOLD
    # run-rate levels
    ws["AF2"], ws["AF3"] = "Run-rate level:", "(last 6 mths / season idx)"
    ws["AF2"].font = BOLD
    ws["AF3"].font = HELP
    for col, act, si in (("AG", "W", "AD"), ("AH", "X", "AE"), ("AI", "Y", "AF"), ("AS", "AQ", "AR")):
        ws[f"{col}2"] = (f"=AVERAGEIFS({RNG(act)},{ACT_WIN('RunRateMonths')})"
                         f"/AVERAGEIFS({RNG(si)},{ACT_WIN('RunRateMonths')})")
        ws[f"{col}2"].number_format = SITE_COLS[col][1]
        ws[f"{col}2"].font = BOLD
    fy_block(ws)
    for i, m in enumerate(MONTHS):
        r = FIRST_ROW + i
        act = f'$D{r}="Actual"'
        ws[f"A{r}"] = m
        ws[f"B{r}"] = f"=YEAR(A{r})"
        ws[f"C{r}"] = f"=MONTH(A{r})"
        ws[f"D{r}"] = f'=IF(A{r}<=LastActual,"Actual","Forecast")'
        if typ == "IDC":
            ws[f"E{r}"] = f"='{host}'!E{r}"
            ws[f"R{r}"] = f'=IF({act},E{r},"")'
        else:
            ws[f"E{r}"] = f"=IF({act},R{r},AB{r}*(1+SalesAdj))"
            ws[f"R{r}"] = f'=IF({act},{data_sum("C", r)},"")'
        ws[f"F{r}"] = f"=IF({act},S{r},AP{r}*AG{r}*(1+CasesAdj))"
        ws[f"AQ{r}"] = f'=IF({act},IFERROR(V{r}/AA{r},""),"")'
        ws[f"AR{r}"] = (f"=IFERROR(AVERAGEIFS({RNG('AQ')},{ACT_WIN('SeasonMonths')},{RNG('C')},$C{r})"
                        f"/AVERAGEIFS({RNG('AQ')},{ACT_WIN('SeasonMonths')}),\"\")")
        ws[f"AS{r}"] = f"=IFERROR($AS$2*AR{r}*AA{r}*(1+TrendLead)^AJ{r},\"\")"
        ws[f"AP{r}"] = ("=" + "+".join(f"'{s}'!E{r}" for s in SERVED[code])) if typ == "IDC" else f"=E{r}"
        ws[f"G{r}"] = f"=IF({act},T{r},IFERROR(H{r}/AH{r},0))"
        ws[f"H{r}"] = f"=IF({act},U{r},J{r}-I{r})"
        ws[f"I{r}"] = f'=IF({act},V{r},IF(LeadMethod="% of sales",E{r}*AI{r},AS{r}))'
        ws[f"J{r}"] = f"=IF({act},U{r}+V{r},E{r}*AC{r})"
        metric_formulas(ws, r)
        for col, src in (("S", "D"), ("T", "E"), ("U", "F"), ("V", "G")):
            ws[f"{col}{r}"] = f'=IF({act},{data_sum(src, r)},"")'
        ws[f"W{r}"] = f'=IF({act},IFERROR(S{r}/AP{r},""),"")'
        ws[f"X{r}"] = f'=IF({act},IFERROR(U{r}/T{r}/Z{r}/AA{r},""),"")'
        ws[f"Y{r}"] = f'=IF({act},IFERROR(V{r}/R{r},""),"")'
        ws[f"Z{r}"] = (f"=INDEX(Calendar!$H$5:$L${4 + N},MATCH($A{r},Calendar!$A$5:$A${4 + N},0),"
                       f"MATCH($D$2,Calendar!$H$4:$L$4,0))")
        ws[f"AA{r}"] = f"=INDEX(Calendar!$M$5:$M${4 + N},MATCH($A{r},Calendar!$A$5:$A${4 + N},0))"
        ws[f"AB{r}"] = f"=SUMIFS({PR('C')},{PR('A')},$A{r},{PR('B')},$B$2)"
        ws[f"AC{r}"] = f"=SUMIFS({PR('D')},{PR('A')},$A{r},{PR('B')},$B$2)"
        for col, src in (("AD", "W"), ("AE", "X"), ("AF", "Y")):
            ws[f"{col}{r}"] = (f"=IFERROR(AVERAGEIFS({RNG(src)},{ACT_WIN('SeasonMonths')},{RNG('C')},$C{r})"
                               f"/AVERAGEIFS({RNG(src)},{ACT_WIN('SeasonMonths')}),\"\")")
        ws[f"AJ{r}"] = f"=MAX(0,(YEAR(A{r})-YEAR(LastActual))*12+MONTH(A{r})-MONTH(LastActual))/12"
        ws[f"AG{r}"] = f"=IFERROR($AG$2*AD{r}*(1+TrendCPS)^AJ{r},\"\")"
        ws[f"AH{r}"] = f"=IFERROR($AH$2*AE{r}*Z{r}*AA{r}*(1+TrendCPH)^AJ{r},\"\")"
        ws[f"AI{r}"] = f"=IFERROR($AI$2*AF{r}*(1+TrendLead)^AJ{r},\"\")"
        # bottom-up: hours needed at LY productivity (+ target), costed at forecast cost per hour
        fc = f'$D{r}="Forecast"'
        ws[f"AK{r}"] = f"=K{r - 12}" if i >= 12 else None
        ws[f"AL{r}"] = f'=IF({fc},IFERROR(F{r}/(AK{r}*(1+ProdTarget)),""),"")'
        ws[f"AM{r}"] = f'=IF({fc},IFERROR(AL{r}*AH{r},""),"")'
        ws[f"AN{r}"] = f'=IF({fc},IFERROR(H{r}-AM{r},""),"")'
        ws[f"AO{r}"] = f'=IF({fc},IFERROR(K{r}/AK{r}-1,""),"")'
    style_site_sheet(ws)


# national first so it sits after Checks; built after the site tabs exist
for s in SITES:
    build_site(*s)

ws = wb.create_sheet("National", index=wb.sheetnames.index("NSW-1"))
ws["A1"] = "National  |  8 regions + 2 IDCs  |  sales exclude IDCs"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "Model status"
ws["B2"] = "=Checks!B3"
ws["B2"].font = BOLD
fy_block(ws)
for i, m in enumerate(MONTHS):
    r = FIRST_ROW + i
    ws[f"A{r}"] = m
    ws[f"B{r}"] = f"=YEAR(A{r})"
    ws[f"C{r}"] = f"=MONTH(A{r})"
    ws[f"D{r}"] = f'=IF(A{r}<=LastActual,"Actual","Forecast")'
    ws[f"E{r}"] = "=" + "+".join(f"'{c}'!E{r}" for c in REGIONS)
    for col in "FGHIJ":
        ws[f"{col}{r}"] = "=" + "+".join(f"'{s[0]}'!{col}{r}" for s in SITES)
    metric_formulas(ws, r)
national_cols = {k: v for k, v in SITE_COLS.items() if k <= "P" and len(k) == 1}
for col, (title, fmt, width) in national_cols.items():
    ws.column_dimensions[col].width = width
    header(ws, 13, ws[f"{col}1"].column, title)
    if col in UNITS:
        u = ws[f"{col}14"]
        u.value, u.font, u.alignment, u.fill = UNITS[col], Font(italic=True), CENTER, LIGHT
ws.row_dimensions[13].height = 45
ws.merge_cells("E12:J12")
header(ws, 12, 5, "Drivers")
ws.merge_cells("K12:P12")
header(ws, 12, 11, "Metrics")
ws.freeze_panes = f"E{FIRST_ROW}"
ws.conditional_formatting.add(f"A{FIRST_ROW}:P{LAST_ROW}", FormulaRule(formula=[f'$D{FIRST_ROW}="Forecast"'], fill=LIGHT))
for r in range(FIRST_ROW, LAST_ROW + 1):
    for col, (_, fmt, _) in national_cols.items():
        if fmt:
            ws[f"{col}{r}"].number_format = fmt
    if MONTHS[r - FIRST_ROW].month == 12:
        for col in "ABCDEFGHIJKLMNOP":
            ws[f"{col}{r}"].border = Border(bottom=THIN)

# site x FY summary tables to the right
ws.column_dimensions["Q"].width = 3
ws.column_dimensions["R"].width = 10
tables = [
    ("Total payroll ($m)", lambda c, y: f"SUMIFS('{c}'!$J${FIRST_ROW}:$J${LAST_ROW},'{c}'!$B${FIRST_ROW}:$B${LAST_ROW},{y})", F_M),
    ("Payroll % of sales (IDC: % of host sales)",
     lambda c, y: (f"IFERROR(SUMIFS('{c}'!$J${FIRST_ROW}:$J${LAST_ROW},'{c}'!$B${FIRST_ROW}:$B${LAST_ROW},{y})"
                   f"/SUMIFS('{c}'!$E${FIRST_ROW}:$E${LAST_ROW},'{c}'!$B${FIRST_ROW}:$B${LAST_ROW},{y}),0)"), F_PCT),
    ("CPLH (cases / hr)",
     lambda c, y: (f"IFERROR(SUMIFS('{c}'!$F${FIRST_ROW}:$F${LAST_ROW},'{c}'!$B${FIRST_ROW}:$B${LAST_ROW},{y})"
                   f"/SUMIFS('{c}'!$G${FIRST_ROW}:$G${LAST_ROW},'{c}'!$B${FIRST_ROW}:$B${LAST_ROW},{y}),0)"), F_1),
    ("Cost per case ($, incl. leadership)",
     lambda c, y: (f"IFERROR(SUMIFS('{c}'!$J${FIRST_ROW}:$J${LAST_ROW},'{c}'!$B${FIRST_ROW}:$B${LAST_ROW},{y})"
                   f"/SUMIFS('{c}'!$F${FIRST_ROW}:$F${LAST_ROW},'{c}'!$B${FIRST_ROW}:$B${LAST_ROW},{y}),0)"), F_CPC),
    ("Staff payroll: top-down minus bottom-up ($m, forecast months; negative = plan looks too tight)",
     lambda c, y: f"SUMIFS('{c}'!$AN${FIRST_ROW}:$AN${LAST_ROW},'{c}'!$B${FIRST_ROW}:$B${LAST_ROW},{y})", F_M),
]
row = 4
for title, fn, fmt in tables:
    ws.cell(row=row, column=18, value=title).font = BOLD
    header(ws, row + 1, 18, "Site")
    for k, y in enumerate(YEARS):
        header(ws, row + 1, 19 + k, f"FY{y}", width=11)
    for i, (code, *_rest) in enumerate(SITES):
        ws.cell(row=row + 2 + i, column=18, value="=" + CODE(code))
        for k, y in enumerate(YEARS):
            c = ws.cell(row=row + 2 + i, column=19 + k, value="=" + fn(code, y))
            c.number_format = fmt
    row += len(SITES) + 4

# ------------------------------------------------------------------ Checks
ws = wb.create_sheet("Checks", index=wb.sheetnames.index("National"))
ws.column_dimensions["A"].width = 62
for col in "BCD":
    ws.column_dimensions[col].width = 16
ws["A1"] = "Checks"
ws["A1"].font = Font(bold=True, size=14)
ws["A3"] = "Overall"
ws["B3"] = '=IF(COUNTIF(D6:D80,"ERROR")=0,"ALL CHECKS OK","CHECK ERRORS")'
ws["B3"].font = BOLD
ws["C3"] = '=COUNTIF(D6:D80,"WARN")&" warning(s)"'
for j, t in enumerate(["Check", "Value", "Expected", "Result"], 1):
    header(ws, 5, j, t)
checks = []
act_months = '(YEAR(LastActual)-2025)*12+MONTH(LastActual)'
fc_months = '(2028-YEAR(LastActual))*12+12-MONTH(LastActual)'
for code, *_r in SITES:
    checks.append((f'={CODE(code)}&": actual rows in Data_Actuals (one per month)"',
                   f'=COUNTIFS({DR("B")},{CODE(code)},{DR("A")},"<="&LastActual)', f"={act_months}", "eq"))
for code, *_r in SITES:
    checks.append((f'={CODE(code)}&": plan payroll % present for every forecast month"',
                   f'=COUNTIFS({PR("B")},{CODE(code)},{PR("A")},">"&LastActual,{PR("D")},">0")',
                   f"={fc_months}", "eq"))
checks.append(("Unknown site codes in Data_Actuals",
               f"=SUMPRODUCT(({DR('B')}<>\"\")*(COUNTIF(SiteCodes,{DR('B')})=0))", "=0", "eq"))
checks.append(("National sales = sum of 8 regions (IDCs excluded)",
               f"=SUM(National!E{FIRST_ROW}:E{LAST_ROW})",
               "=" + "+".join(f"SUM('{c}'!E{FIRST_ROW}:E{LAST_ROW})" for c in REGIONS), "eq"))
checks.append(("National total payroll = sum of all 10 sites",
               f"=SUM(National!J{FIRST_ROW}:J{LAST_ROW})",
               "=" + "+".join(f"SUM('{s[0]}'!J{FIRST_ROW}:J{LAST_ROW})" for s in SITES), "eq"))
checks.append(("Staff + leadership = total payroll (national, all months)",
               f"=SUM(National!H{FIRST_ROW}:I{LAST_ROW})", f"=SUM(National!J{FIRST_ROW}:J{LAST_ROW})", "eq"))
checks.append(("Months with negative staff payroll (leadership > total)",
               "=" + "+".join(f"COUNTIF('{s[0]}'!H{FIRST_ROW}:H{LAST_ROW},\"<0\")" for s in SITES), "=0", "eq"))
checks.append(("At least 12 actual months for seasonality", f"={act_months}", "=SeasonMonths", "ge"))
for code, *_r in SITES:
    rng = f"'{code}'!AO{FIRST_ROW}:AO{LAST_ROW}"
    checks.append((f'={CODE(code)}&": forecast months where CPLH moves more than the tolerance vs LY"',
                   f'=COUNTIF({rng},">"&CPLHTol)+COUNTIF({rng},"<"&-CPLHTol)', "=0", "warn"))
for i, (label, val, exp, op) in enumerate(checks):
    r = 6 + i
    ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = label, val, exp
    test = f"B{r}>=C{r}" if op == "ge" else f"ABS(B{r}-C{r})<=MAX(1,ABS(C{r})*0.000001)"
    ws[f"D{r}"] = f'=IF({test},"OK","{"WARN" if op == "warn" else "ERROR"}")'
    for col in "BC":
        ws[f"{col}{r}"].number_format = F_INT
ws.conditional_formatting.add("D6:D80", FormulaRule(formula=['D6="ERROR"'], fill=PatternFill("solid", fgColor="F8CBAD")))
ws.conditional_formatting.add("D6:D80", FormulaRule(formula=['D6="OK"'], fill=PatternFill("solid", fgColor="C6EFCE")))
ws.conditional_formatting.add("D6:D80", FormulaRule(formula=['D6="WARN"'], fill=PatternFill("solid", fgColor="FFEB9C")))

# ---------------------------------------------------------------- Backtest
ws = wb.create_sheet("Backtest", index=wb.sheetnames.index("National"))
ws["A1"] = "Backtest: forecast vs actual"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = ("To backtest, set 'Last actual month' on Inputs earlier (e.g. Mar-26). Months after it that have actuals "
            "in Data_Actuals are compared with the model. Set it back afterwards.")
ws["A2"].font = HELP
BT = [("Cases", "D", "F", F_CASES), ("Hours", "E", "G", F_INT), ("Staff payroll", "F", "H", F_M),
      ("Leadership", "G", "I", F_M), ("CPLH", None, "K", F_1)]
ws["A4"] = "Mean absolute error"
ws["A4"].font = BOLD
header(ws, 7, 1, "Month", width=9)
header(ws, 7, 2, "Status", width=9)
for k, (title, dcol, mcol, fmt) in enumerate(BT):
    c = 3 + 3 * k
    ws.merge_cells(start_row=6, start_column=c, end_row=6, end_column=c + 2)
    header(ws, 6, c, title)
    for j, t in enumerate(["Actual", "Model", "Error"]):
        header(ws, 7, c + j, t, width=12)
    e = get_column_letter(c + 2)
    ws.cell(row=4, column=c + 2, value=(f'=IFERROR((SUMIF({e}8:{e}{7 + N},">0")-SUMIF({e}8:{e}{7 + N},"<0"))'
                                         f'/COUNT({e}8:{e}{7 + N}),"n/a")')).number_format = F_PCT
    ws.cell(row=4, column=c + 2).font = BOLD
    for i, m in enumerate(MONTHS):
        r, src = 8 + i, FIRST_ROW + i
        a, mo, er = (get_column_letter(c + j) for j in range(3))
        if dcol:
            ws[f"{a}{r}"] = f"=SUMIFS({DR(dcol)},{DR('A')},$A{r})"
        else:  # CPLH = actual cases / actual hours
            ws[f"{a}{r}"] = f"=IFERROR(C{r}/F{r},0)"
        ws[f"{mo}{r}"] = f"=National!{mcol}{src}"
        ws[f"{er}{r}"] = f'=IF(AND($B{r}="Forecast",{a}{r}>0),{mo}{r}/{a}{r}-1,"")'
        ws[f"{a}{r}"].number_format = ws[f"{mo}{r}"].number_format = fmt
        ws[f"{er}{r}"].number_format = F_PCT
for i, m in enumerate(MONTHS):
    r = 8 + i
    ws[f"A{r}"] = m
    ws[f"A{r}"].number_format = F_DATE
    ws[f"B{r}"] = f"=National!D{FIRST_ROW + i}"
ws.freeze_panes = "C8"

# --------------------------------------------------------------- Dashboard
SERIES_ACT, SERIES_FC = "2A78D6", "EB6834"  # categorical slots 1 and 2
ws = wb.create_sheet("Dashboard", index=wb.sheetnames.index("National"))
ws["A1"] = "Dashboard"
ws["A1"].font = Font(bold=True, size=14)
ws["A3"] = "Site"
ws["B3"] = "National"
ws["B3"].fill, ws["B3"].font = INPUT, Font(bold=True, color="0000FF")
site_dv = DataValidation(type="list", formula1=f"$Y$4:$Y${4 + len(SITES)}")
ws.add_data_validation(site_dv)
site_dv.add("B3")
ws["C3"] = "Pick a site to update every chart. Blue = actual, orange = forecast."
ws["C3"].font = HELP
ws["A4"] = "Status"
ws["B4"] = "=Checks!B3"
DASH_SITES = ["National"] + [c for c, *_ in SITES]
for i, nm in enumerate(DASH_SITES):
    ws.cell(row=4 + i, column=25, value=nm if i == 0 else "=" + CODE(nm)).font = HELP  # column Y: selector list
ws["X3"] = "Site no."
ws["X4"] = f"=MATCH($B$3,$Y$4:$Y${3 + len(DASH_SITES)},0)"
ws["X3"].font = ws["X4"].font = HELP

# chart data for the selected site (helper block from column AA)
METRICS = [("Total payroll", "J", '#,##0.0,,"m"'), ("WH hours worked", "G", '#,##0'),
           ("CPLH (cases / hr)", "K", '0'), ("Cost per hour", "L", '$0'),
           ("Cost per case (incl. leadership)", "M", '$0.00'), ("Total payroll % of sales", "P", '0.0%')]
C0 = 27  # column AA
ws.cell(row=2, column=C0, value="Chart data for the selected site (do not edit)").font = HELP
ws.cell(row=3, column=C0, value="Month").font = BOLD
for k, (title, col, fmt) in enumerate(METRICS):
    ws.cell(row=3, column=C0 + 1 + 2 * k, value=f"{title}: actual").font = BOLD
    ws.cell(row=3, column=C0 + 2 + 2 * k, value=f"{title}: forecast").font = BOLD
for i, m in enumerate(MONTHS):
    r, src = 4 + i, FIRST_ROW + i
    ws.cell(row=r, column=C0, value=m).number_format = F_DATE
    pick = lambda col: "CHOOSE($X$4," + ",".join(f"'{nm}'!{col}{src}" for nm in DASH_SITES) + ")"
    status = pick("D")
    for k, (title, col, fmt) in enumerate(METRICS):
        val = pick(col)
        a = ws.cell(row=r, column=C0 + 1 + 2 * k, value=f'=IF({status}="Actual",{val},NA())')
        # forecast line starts at the last actual month so the two lines join
        f = ws.cell(row=r, column=C0 + 2 + 2 * k,
                    value=f'=IF(OR({status}="Forecast",${get_column_letter(C0)}{r}=LastActual),{val},NA())')
        a.number_format = f.number_format = fmt
D_LAST = 3 + N


def line_chart(title, k, fmt):
    ch = LineChart()
    ch.title = title
    ch.height, ch.width = 7.5, 16
    ch.legend.position = "b"
    ch.y_axis.number_format = fmt
    ch.y_axis.majorGridlines.spPr = None
    ch.x_axis.number_format = "mmm-yy"
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.x_axis.tickLblSkip = 6
    for off, color in ((1, SERIES_ACT), (2, SERIES_FC)):
        ref = Reference(ws, min_col=C0 + off + 2 * k, min_row=3, max_row=D_LAST)
        ch.add_data(ref, titles_from_data=True)
        ser = ch.series[-1]
        ser.graphicalProperties.line.solidFill = color
        ser.graphicalProperties.line.width = 25400  # 2pt
        ser.smooth = False
        ser.marker.symbol = "none"
        ser.tx.strRef = None
        from openpyxl.chart.series import SeriesLabel
        ser.tx = SeriesLabel(v="Actual" if off == 1 else "Forecast")
    ch.set_categories(Reference(ws, min_col=C0, min_row=4, max_row=D_LAST))
    return ch


for k, (title, col, fmt) in enumerate(METRICS):
    anchor = f"{'A' if k % 2 == 0 else 'L'}{6 + (k // 2) * 16}"
    ws.add_chart(line_chart(title, k, fmt), anchor)

# gap chart: top-down minus bottom-up staff payroll by site, FY2027 and FY2028 (from the National summary)
gap_title_row = 4 + 4 * (len(SITES) + 4)
bar = BarChart()
bar.type = "bar"
bar.title = "Plan payroll minus bottom-up need, by site (negative = plan looks too tight)"
bar.height, bar.width = 9, 33
bar.y_axis.number_format = '#,##0.0,,"m"'
bar.x_axis.delete = bar.y_axis.delete = False
bar.y_axis.majorGridlines.spPr = None
bar.legend.position = "b"
bar.gapWidth = 60
nat = wb["National"]
for col, color in ((21, SERIES_ACT), (22, SERIES_FC)):
    bar.add_data(Reference(nat, min_col=col, min_row=gap_title_row + 1, max_row=gap_title_row + 1 + len(SITES)),
                 titles_from_data=True)
    bar.series[-1].graphicalProperties.solidFill = color
    bar.series[-1].graphicalProperties.line.noFill = True
bar.set_categories(Reference(nat, min_col=18, min_row=gap_title_row + 2, max_row=gap_title_row + 1 + len(SITES)))
ws.add_chart(bar, f"A{6 + 3 * 16}")
for col in "ABCDEFGHIJKLMNOPQRSTUVW":
    ws.column_dimensions[col].width = 9

# tab colours
for nm in wb.sheetnames:
    if nm in ("Inputs", "Data_Actuals", "Plan_Inputs", "Holidays"):
        wb[nm].sheet_properties.tabColor = "FFC000"
    elif nm in REGIONS:
        wb[nm].sheet_properties.tabColor = "1F3864"
    elif nm in ("IDC-NSW", "IDC-VIC"):
        wb[nm].sheet_properties.tabColor = "2E75B6"
    elif nm in ("National", "Dashboard", "Backtest"):
        wb[nm].sheet_properties.tabColor = "00B050"

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
